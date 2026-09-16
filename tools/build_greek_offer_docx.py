#!/usr/bin/env python3
"""Helikon Consortium — Word copy of the one-page final offer."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "Helikon-Consortium-Offer-Jollof-Living.docx")


def set_run_font(run, name="Times New Roman", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "222222")
    pBdr.append(bottom)
    pPr.append(pBdr)


def p(doc, text, size=11, bold=False, italic=False, color=None, space_after=8, space_before=0, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing = 1.15
    if align:
        para.alignment = align
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return para


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    ink = RGBColor(0x11, 0x11, 0x11)
    muted = RGBColor(0x44, 0x44, 0x44)

    head = doc.add_paragraph()
    head.paragraph_format.space_after = Pt(2)
    r = head.add_run("HELIKON CONSORTIUM")
    set_run_font(r, size=14, bold=True, color=ink)
    r2 = head.add_run("\tAthens, Hellenic Republic")
    set_run_font(r2, size=10, color=muted)
    tab = head.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_ALIGN_PARAGRAPH.RIGHT)
    add_bottom_border(head)

    addr = doc.add_paragraph()
    addr.paragraph_format.space_before = Pt(8)
    addr.paragraph_format.space_after = Pt(14)
    a1 = addr.add_run("12 Leoforos Vasilissis Sofias, 106 71 Athens")
    set_run_font(a1, size=10, color=muted)
    a2 = addr.add_run("\t16 September 2026")
    set_run_font(a2, size=10, color=muted)

    p(doc, "PRIVATE & CONFIDENTIAL", size=11, bold=True, color=ink, space_after=10)
    p(doc, "Mr Chinwike Ogbunugafor and Co.", size=11, color=ink, space_after=0)
    p(doc, "100% owners of the company", size=11, color=muted, space_after=14)
    p(doc, "Dear Mr Ogbunugafor,", size=11, color=ink, space_after=10)

    body = [
        "Following our meetings, Helikon Consortium hereby offers to buy ninety per cent (90%) of the entire issued share capital of the company from you and your co-owners.",
        "Price: ₦2.88 billion (US$1.80 million), being eighteen per cent (18%) of the Year-3 enterprise value of ₦16 billion in the financial projections we reviewed. Cash at closing, in USD or NGN. You retain ten per cent (10%) as ordinary shares. No earn-out, no further cash.",
        "Mr Chinwike Ogbunugafor will serve as Chief Technology Officer for eighteen (18) months from closing at ₦500,000 per month. No bonus. No additional equity.",
        "This offer is final and absolute. It will not be negotiated further. A counter-proposal is a rejection. Accept by signing below and returning this letter by 17:00 WAT on 30 September 2026, failing which the offer lapses. Closing by 31 October 2026. Nigerian law; Lagos courts. We have the funds; there is no financing condition.",
    ]
    for t in body:
        p(doc, t, size=11, color=ink, space_after=10)

    p(doc, "Yours faithfully,", size=11, color=ink, space_after=16, space_before=4)
    p(doc, "Helikon Consortium", size=11, bold=True, color=ink, space_after=0)
    p(doc, "Authorised signatories", size=10, color=muted, space_after=18)

    acc_h = doc.add_paragraph()
    acc_h.paragraph_format.space_before = Pt(6)
    acc_h.paragraph_format.space_after = Pt(8)
    add_bottom_border(acc_h)
    rh = acc_h.add_run("ACCEPTANCE")
    set_run_font(rh, size=11, bold=True, color=ink)

    p(
        doc,
        "We, being the 100% owners, accept Helikon Consortium’s offer of 16 September 2026 to purchase 90% of the company for ₦2.88 billion (US$1.80 million), we retaining 10%, with Mr Chinwike Ogbunugafor as CTO for 18 months at ₦500,000 per month. We acknowledge the offer is final and will not be negotiated further.",
        size=10,
        color=ink,
        space_after=22,
    )

    sig = doc.add_table(rows=2, cols=3)
    sig.autofit = True
    labels = ["Mr Chinwike Ogbunugafor", "Co-owner", "Co-owner"]
    for i, lab in enumerate(labels):
        cell0 = sig.cell(0, i)
        cell0.text = "________________________"
        for run in cell0.paragraphs[0].runs:
            set_run_font(run, size=10, color=ink)
        cell1 = sig.cell(1, i)
        cell1.paragraphs[0].clear()
        r = cell1.paragraphs[0].add_run(lab + "  ·  date")
        set_run_font(r, size=9, color=muted)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("Wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    build()
